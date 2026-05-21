import io
from datetime import date
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.models import SDRContext
from utils.formatting import fmt_currency, fmt_pct


# ── color palette ────────────────────────────────────────────────
AZUL_ESCURO = RGBColor(0, 59, 92)
AZUL_MEDIO = RGBColor(0, 102, 153)
VERDE = RGBColor(0, 128, 64)
VERMELHO = RGBColor(192, 0, 0)
AMARELO = RGBColor(255, 192, 0)
LARANJA = RGBColor(255, 102, 0)
CINZA_CLARO = RGBColor(242, 242, 242)
BRANCO = RGBColor(255, 255, 255)


def _set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_text(cell, text: str, bold=False, color: RGBColor = None, size=10,
                   align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _add_table_from_df(doc: Document, df: pd.DataFrame,
                        header_bg: RGBColor = None,
                        limbo_col: str = None) -> None:
    header_bg = header_bg or AZUL_MEDIO
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"

    # Header row
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        _set_cell_bg(hdr_cells[i], header_bg)
        _set_cell_text(hdr_cells[i], col, bold=True, color=BRANCO,
                       align=WD_ALIGN_PARAGRAPH.CENTER)

    # Data rows
    for row_idx, row in df.iterrows():
        row_cells = table.add_row().cells
        bg = CINZA_CLARO if row_idx % 2 == 1 else BRANCO
        for col_idx, (col, val) in enumerate(zip(df.columns, row)):
            cell = row_cells[col_idx]
            if limbo_col and col == limbo_col and str(val).strip() == "Limbo":
                _set_cell_bg(cell, VERMELHO)
                _set_cell_text(cell, val, bold=True, color=BRANCO)
            else:
                _set_cell_bg(cell, bg)
                _set_cell_text(cell, val if val is not None else "")


def _add_kpi_table(doc: Document, kpis: list) -> None:
    """kpis = [(label, value, color), ...]"""
    table = doc.add_table(rows=2, cols=len(kpis))
    table.style = "Table Grid"
    for i, (label, value, color) in enumerate(kpis):
        _set_cell_bg(table.rows[0].cells[i], AZUL_ESCURO)
        _set_cell_text(table.rows[0].cells[i], label, bold=True, color=BRANCO,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_bg(table.rows[1].cells[i], color or CINZA_CLARO)
        _set_cell_text(table.rows[1].cells[i], value, bold=True, size=12,
                       align=WD_ALIGN_PARAGRAPH.CENTER)


def gerar_word(ctx: SDRContext, r1: dict, r2: dict, r3: dict, r4: dict) -> io.BytesIO:
    doc = Document()

    # ── Margens ─────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── CAPA ────────────────────────────────────────────────────
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ANÁLISE DE FORECAST SDR")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = AZUL_ESCURO

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"SDR: {ctx.nome}  |  {ctx.senioridade}  |  {ctx.mes_referencia}").font.size = Pt(14)

    doc.add_paragraph()
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"Meta NR: {fmt_currency(ctx.meta_net_revenue)}  |  Meta SQLs: {ctx.meta_sqls}"
    ).font.size = Pt(12)

    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.add_run(f"Gerado em: {date.today().strftime('%d/%m/%Y')}").font.size = Pt(10)

    doc.add_page_break()

    # ── SEÇÃO 1 — FORECAST DO MÊS ───────────────────────────────
    doc.add_heading("1. Forecast do Mês", level=1)

    ating_pct = r1["ating_ponderado"] * 100
    color_ating = VERDE if ating_pct >= 80 else (AMARELO if ating_pct >= 50 else VERMELHO)

    _add_kpi_table(doc, [
        ("NR Projetado", fmt_currency(r1["nr_projetado"]), CINZA_CLARO),
        ("Pipeline Ativação", fmt_currency(r1["nr_pipeline"]), CINZA_CLARO),
        ("SQLs Quentes", str(r1["sqls_quentes"]), CINZA_CLARO),
        ("Atingimento Ponderado", f"{ating_pct:.1f}%", color_ating),
    ])

    doc.add_paragraph()
    doc.add_heading("1.1 Carteira Ativa por Mês de Transição", level=2)
    resumo = r1["df_resumo_por_mes"].copy()
    resumo.columns = ["Mês na Carteira", "Qtd Clientes", "NR Total (R$)"]
    resumo["NR Total (R$)"] = resumo["NR Total (R$)"].apply(fmt_currency)
    _add_table_from_df(doc, resumo)

    doc.add_paragraph()
    doc.add_heading("1.2 Saídas Obrigatórias (Mês 3)", level=2)
    saidas = r1["df_saidas"]
    if saidas.empty:
        doc.add_paragraph("Nenhuma saída obrigatória neste mês.")
    else:
        cols = ["Account Name", "Net Revenue", "Months from Activation"]
        cols = [c for c in cols if c in saidas.columns]
        s = saidas[cols].copy()
        if "Net Revenue" in s.columns:
            s["Net Revenue"] = s["Net Revenue"].apply(fmt_currency)
        _add_table_from_df(doc, s, header_bg=VERMELHO)

    doc.add_paragraph()
    doc.add_heading("1.3 Pipeline de Ativação", level=2)
    pipe = r1["df_pipeline_ativacao"]
    if pipe.empty:
        doc.add_paragraph("Nenhum cliente no pipeline de ativação.")
    else:
        cols = ["Account Name", "Amount", "GMV Faltante", "Probabilidade", "NR Esperado"]
        cols = [c for c in cols if c in pipe.columns]
        p = pipe[cols].copy()
        if "Amount" in p.columns:
            p["Amount"] = p["Amount"].apply(fmt_currency)
        if "GMV Faltante" in p.columns:
            p["GMV Faltante"] = p["GMV Faltante"].apply(fmt_currency)
        if "Probabilidade" in p.columns:
            p["Probabilidade"] = p["Probabilidade"].apply(fmt_pct)
        if "NR Esperado" in p.columns:
            p["NR Esperado"] = p["NR Esperado"].apply(fmt_currency)
        _add_table_from_df(doc, p)

    doc.add_page_break()

    # ── SEÇÃO 2 — SEGMENTAÇÃO N1-N7 ─────────────────────────────
    doc.add_heading("2. Segmentação de Oportunidades por Nível", level=1)

    tabela_resumo = r2["tabela_resumo"].copy()
    tabela_resumo["Amount Total (R$)"] = tabela_resumo["Amount Total (R$)"].apply(fmt_currency)
    tabela_resumo["Amount Won (R$)"] = tabela_resumo["Amount Won (R$)"].apply(fmt_currency)
    _add_table_from_df(doc, tabela_resumo)

    for nivel, df_nivel in r2["detalhe_por_nivel"].items():
        if df_nivel.empty:
            continue
        doc.add_paragraph()
        doc.add_heading(f"2.{nivel} — Oportunidades {nivel}", level=2)
        d = df_nivel.copy()
        if "Amount" in d.columns:
            d["Amount"] = d["Amount"].apply(fmt_currency)
        if "Close Date" in d.columns:
            d["Close Date"] = d["Close Date"].apply(
                lambda x: x.strftime("%d/%m/%Y") if pd.notna(x) else ""
            )
        _add_table_from_df(doc, d)

    doc.add_page_break()

    # ── SEÇÃO 3 — CRUZAMENTO NÍVEL × STATUS ─────────────────────
    doc.add_heading("3. Cruzamento Nível × Status da Carteira", level=1)

    pivot_count = r3["pivot_count"].reset_index()
    _add_table_from_df(doc, pivot_count)

    doc.add_paragraph()
    doc.add_heading("3.1 Amount por Nível × Status (R$)", level=2)
    pivot_amount = r3["pivot_amount"].reset_index().copy()
    for col in pivot_amount.columns[1:]:
        pivot_amount[col] = pivot_amount[col].apply(fmt_currency)
    _add_table_from_df(doc, pivot_amount)

    limbo = r3["limbo_detail"]
    limbo_non_zero = limbo[limbo["amount_total"] > 0] if not limbo.empty else limbo
    if not limbo_non_zero.empty:
        doc.add_paragraph()
        doc.add_heading("3.2 Contas em Limbo — Ação Recomendada", level=2)
        p = doc.add_paragraph()
        run = p.add_run(
            f"⚠ Total em Limbo: {fmt_currency(r3['total_limbo_amount'])}"
        )
        run.bold = True
        run.font.color.rgb = VERMELHO
        l = limbo_non_zero.copy()
        l.columns = ["Nível", "Qtd Contas", "Amount Total (R$)"]
        l["Amount Total (R$)"] = l["Amount Total (R$)"].apply(fmt_currency)
        _add_table_from_df(doc, l, header_bg=VERMELHO)

    doc.add_page_break()

    # ── SEÇÃO 4 — SIMULAÇÃO SQLS N3/N4 ──────────────────────────
    doc.add_heading("4. Simulação de Meta de SQLs N3 e N4", level=1)

    doc.add_heading("4.1 Taxas de Conversão Históricas por Nível", level=2)
    taxas = r4["df_taxas_conversao"].copy()
    taxas["Taxa SQL→Won"] = taxas["Taxa SQL→Won"].apply(fmt_pct)
    taxas["Ticket Médio Won (R$)"] = taxas["Ticket Médio Won (R$)"].apply(fmt_currency)
    _add_table_from_df(doc, taxas)

    doc.add_paragraph()
    doc.add_heading("4.2 Composição Ideal da Carteira (17 clientes)", level=2)
    comp = r4["df_composicao_ideal"].copy()
    comp["Ticket Médio (R$)"] = comp["Ticket Médio (R$)"].apply(fmt_currency)
    comp["NR Projetado (R$)"] = comp["NR Projetado (R$)"].apply(fmt_currency)
    _add_table_from_df(doc, comp)

    doc.add_paragraph()
    doc.add_heading("4.3 SQLs Necessários por Mês", level=2)
    dist = r4["df_distribuicao_sqls"].copy()
    dist["Taxa SQL→Won"] = dist["Taxa SQL→Won"].apply(fmt_pct)
    _add_table_from_df(doc, dist)

    sqls_n3n4 = r4["sqls_necessarios_n3_n4"]
    p = doc.add_paragraph()
    p.add_run("Recomendação: ").bold = True
    n3 = sqls_n3n4.get("N3")
    n4 = sqls_n3n4.get("N4")
    p.add_run(
        f"Gerar ~{n3 or 'N/D'} SQLs N3/mês e ~{n4 or 'N/D'} SQLs N4/mês "
        "para sustentar a carteira ideal."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
